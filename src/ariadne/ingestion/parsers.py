"""Extracao de texto de formatos variados.

Cada formato quebra de um jeito diferente, e tratar todos como "abrir e ler"
produz corpus ruim sem avisar:

- PDF nativo tem texto embutido; PDF ESCANEADO e so imagem, e devolve zero
  caractere em silencio. O parser detecta esse caso e cai no OCR em vez de
  indexar um documento vazio.
- DOCX guarda titulos como estilo, nao como marcacao. Convertidos para `==`,
  eles alimentam o chunking por secao -- sem isso o documento inteiro vira um
  bloco sem hierarquia.
- Planilha nao e texto corrido. Chunkar linhas de tabela destroi o
  significado: "R$ 1.200" sozinho nao diz nada. Cada linha vira uma FRASE com
  os nomes das colunas, que e o que o embedding consegue representar.
- Imagem so tem pixel. Vai para OCR, e o resultado vem marcado como tal --
  texto de OCR erra, e quem le a resposta merece saber disso.

Todas as dependencias sao opcionais: a ausencia de uma vira mensagem util, nao
ImportError no meio da ingestao.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from pathlib import Path
from typing import Any

PDF_SUFFIXES = {".pdf"}
DOCX_SUFFIXES = {".docx"}
SHEET_SUFFIXES = {".xlsx", ".xlsm"}
CSV_SUFFIXES = {".csv", ".tsv"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff", ".tif"}
TEXT_SUFFIXES = {".md", ".markdown", ".txt", ".rst"}

SUPPORTED = (
    TEXT_SUFFIXES | PDF_SUFFIXES | DOCX_SUFFIXES | SHEET_SUFFIXES | CSV_SUFFIXES | IMAGE_SUFFIXES
)

# Abaixo disso, um PDF de N paginas provavelmente e digitalizacao.
_MIN_CHARS_POR_PAGINA = 40


class ParserUnavailableError(RuntimeError):
    """Formato reconhecido, mas a dependencia nao esta instalada."""


@dataclass
class ParsedDocument:
    text: str
    kind: str
    """Como o texto foi obtido: 'texto', 'pdf', 'pdf-ocr', 'docx', 'planilha', 'ocr'."""

    pages: int = 0
    warning: str | None = None


def parse_file(path: Path) -> ParsedDocument:
    """Extrai texto de um arquivo, escolhendo a estrategia pela extensao."""
    sufixo = path.suffix.lower()

    if sufixo in TEXT_SUFFIXES:
        return ParsedDocument(text=path.read_text(encoding="utf-8", errors="replace"), kind="texto")
    if sufixo in PDF_SUFFIXES:
        return _parse_pdf(path)
    if sufixo in DOCX_SUFFIXES:
        return _parse_docx(path)
    if sufixo in SHEET_SUFFIXES:
        return _parse_sheet(path)
    if sufixo in CSV_SUFFIXES:
        return _parse_csv(path)
    if sufixo in IMAGE_SUFFIXES:
        return _parse_image(path)

    msg = f"extensao {sufixo!r} nao suportada"
    raise ParserUnavailableError(msg)


def _parse_pdf(path: Path) -> ParsedDocument:
    try:
        import pymupdf
        import pymupdf4llm  # type: ignore[import-untyped]
    except ImportError as exc:  # pragma: no cover
        msg = "instale o extra 'docs' para ler PDF: uv sync --extra docs"
        raise ParserUnavailableError(msg) from exc

    documento: Any = pymupdf.open(path)  # type: ignore[no-untyped-call]
    try:
        paginas = int(documento.page_count)
    finally:
        documento.close()

    texto: str = pymupdf4llm.to_markdown(str(path), show_progress=False)

    # PDF escaneado devolve quase nada e NAO levanta erro -- indexar assim
    # criaria um documento vazio que some na busca sem ninguem entender por que.
    if paginas and len(texto.strip()) < paginas * _MIN_CHARS_POR_PAGINA:
        ocr = _ocr_pdf(path)
        if ocr and len(ocr.strip()) > len(texto.strip()):
            return ParsedDocument(
                text=ocr,
                kind="pdf-ocr",
                pages=paginas,
                warning=(
                    "PDF sem texto embutido (provavelmente digitalizado); "
                    "o conteudo veio de OCR e pode conter erros de leitura"
                ),
            )

    return ParsedDocument(text=texto, kind="pdf", pages=paginas)


def _parse_docx(path: Path) -> ParsedDocument:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        msg = "instale o extra 'docs' para ler .docx: uv sync --extra docs"
        raise ParserUnavailableError(msg) from exc

    documento: Any = docx.Document(str(path))
    linhas: list[str] = []
    for paragrafo in documento.paragraphs:
        texto = paragrafo.text.strip()
        if not texto:
            continue
        estilo = ((paragrafo.style.name if paragrafo.style else "") or "").lower()
        if estilo.startswith(("heading", "titulo", "título")):
            # Converte o titulo para a marcacao que o chunking entende, para o
            # documento manter hierarquia em vez de virar um bloco unico.
            nivel = "".join(c for c in estilo if c.isdigit()) or "1"
            marcas = "=" * (min(int(nivel), 4) + 1)
            linhas.append(f"\n{marcas} {texto} {marcas}\n")
        else:
            linhas.append(texto)

    # Tabelas viram linhas legiveis, nao grade.
    for tabela in documento.tables:
        for linha in tabela.rows:
            celulas = [c.text.strip() for c in linha.cells if c.text.strip()]
            if celulas:
                linhas.append(" | ".join(celulas))

    return ParsedDocument(text="\n\n".join(linhas), kind="docx")


def _parse_sheet(path: Path) -> ParsedDocument:
    try:
        import openpyxl  # type: ignore[import-untyped]
    except ImportError as exc:  # pragma: no cover
        msg = "instale o extra 'docs' para ler planilhas: uv sync --extra docs"
        raise ParserUnavailableError(msg) from exc

    livro = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    partes: list[str] = []
    for aba in livro.worksheets:
        linhas = list(aba.iter_rows(values_only=True))
        if not linhas:
            continue
        partes.append(f"\n== {aba.title} ==\n")
        partes.extend(_linhas_para_frases(linhas))
    livro.close()
    return ParsedDocument(text="\n".join(partes), kind="planilha")


def _parse_csv(path: Path) -> ParsedDocument:
    bruto = path.read_text(encoding="utf-8", errors="replace")
    delimitador = "\t" if path.suffix.lower() == ".tsv" else _detectar_delimitador(bruto)
    linhas = list(csv.reader(io.StringIO(bruto), delimiter=delimitador))
    return ParsedDocument(text="\n".join(_linhas_para_frases(linhas)), kind="planilha")


def _linhas_para_frases(linhas: list[Any]) -> list[str]:
    """Converte linhas de tabela em frases com o nome de cada coluna.

    "R$ 1.200" isolado nao significa nada para um embedding. "Valor: R$ 1.200;
    Fornecedor: Alfa" significa. E a diferenca entre a planilha ser
    pesquisavel e virar ruido numerico no indice.
    """
    if not linhas:
        return []
    cabecalho = [str(c).strip() if c is not None else "" for c in linhas[0]]
    frases: list[str] = []
    for linha in linhas[1:]:
        campos = [
            f"{cabecalho[i] or f'coluna {i + 1}'}: {valor}"
            for i, valor in enumerate(linha)
            if i < len(cabecalho) and valor is not None and str(valor).strip()
        ]
        if campos:
            frases.append("; ".join(campos))
    return frases


def _detectar_delimitador(texto: str) -> str:
    amostra = texto[:4096]
    try:
        return csv.Sniffer().sniff(amostra, delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def _parse_image(path: Path) -> ParsedDocument:
    texto = _ocr_image(path)
    return ParsedDocument(
        text=texto,
        kind="ocr",
        warning=(
            "conteudo extraido por OCR: erros de leitura sao esperados, "
            "principalmente em foto de celular ou letra manuscrita"
        ),
    )


def _ocr_engine() -> Any:
    try:
        from rapidocr_onnxruntime import (  # type: ignore[import-untyped]
            RapidOCR,
        )
    except ImportError as exc:  # pragma: no cover
        msg = "instale o extra 'ocr' para ler imagens: uv sync --extra ocr"
        raise ParserUnavailableError(msg) from exc
    return RapidOCR()


def _ocr_image(path: Path) -> str:
    motor = _ocr_engine()
    resultado, _ = motor(str(path))
    if not resultado:
        return ""
    return "\n".join(linha[1] for linha in resultado if len(linha) > 1)


def _ocr_pdf(path: Path) -> str:
    """Renderiza cada pagina como imagem e passa no OCR."""
    try:
        import pymupdf
    except ImportError:  # pragma: no cover
        return ""

    try:
        motor = _ocr_engine()
    except ParserUnavailableError:
        return ""

    partes: list[str] = []
    documento: Any = pymupdf.open(path)  # type: ignore[no-untyped-call]
    try:
        for pagina in documento:
            # 200 dpi: abaixo disso o OCR erra muito, acima fica lento sem
            # ganho perceptivel em documento de texto.
            pix = pagina.get_pixmap(dpi=200)
            resultado, _ = motor(pix.tobytes("png"))
            if resultado:
                partes.append("\n".join(linha[1] for linha in resultado if len(linha) > 1))
    finally:
        documento.close()
    return "\n\n".join(partes)
